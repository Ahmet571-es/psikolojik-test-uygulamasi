# -*- coding: utf-8 -*-
"""Eğitim Check-Up - Aile Bilgilendirme Rapor Şablonu (PDF/DOCX)"""

import io
import os
from datetime import datetime


def _font_yolu_bul():
    """DejaVu Sans font dosyasını bul"""
    yollar = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/dejavu/DejaVuSans.ttf",
        os.path.join(os.path.dirname(__file__), "fonts", "DejaVuSans.ttf"),
    ]
    for yol in yollar:
        if os.path.exists(yol):
            return yol
    return None


def aile_pdf_olustur(ogrenci_adi, ozet_metni, ogretmen_adi="", secilen_basliklar=None):
    """Aile bilgilendirme PDF raporu oluştur"""
    try:
        from fpdf import FPDF
        from fpdf.enums import XPos, YPos

        font_yolu = _font_yolu_bul()

        pdf = FPDF()
        font_adi = "Helvetica"

        if font_yolu and os.path.exists(font_yolu):
            pdf.add_font("DejaVu", "", font_yolu)
            bold_yol = font_yolu.replace("DejaVuSans.ttf", "DejaVuSans-Bold.ttf")
            if os.path.exists(bold_yol):
                pdf.add_font("DejaVu", "B", bold_yol)
            else:
                pdf.add_font("DejaVu", "B", font_yolu)
            font_adi = "DejaVu"

        pdf.set_auto_page_break(auto=True, margin=25)
        pdf.add_page()

        # Başlık
        pdf.set_draw_color(59, 130, 246)
        pdf.set_line_width(1)
        pdf.line(10, 10, 200, 10)
        pdf.ln(8)

        pdf.set_font(font_adi, "B", 18)
        pdf.set_text_color(30, 41, 59)
        pdf.cell(0, 12, "Veli Bilgilendirme Özeti",
                 new_x=XPos.LMARGIN, new_y=YPos.NEXT, align="C")

        pdf.set_font(font_adi, "", 10)
        pdf.set_text_color(100, 116, 139)
        pdf.cell(0, 6, "Eğitim Check-Up Platformu",
                 new_x=XPos.LMARGIN, new_y=YPos.NEXT, align="C")

        pdf.set_draw_color(226, 232, 240)
        pdf.set_line_width(0.3)
        pdf.line(10, pdf.get_y() + 2, 200, pdf.get_y() + 2)
        pdf.ln(8)

        # Öğrenci ve tarih bilgileri
        pdf.set_font(font_adi, "B", 11)
        pdf.set_text_color(30, 41, 59)
        pdf.cell(0, 8, f"Öğrenci: {ogrenci_adi}",
                 new_x=XPos.LMARGIN, new_y=YPos.NEXT)

        pdf.set_font(font_adi, "", 10)
        pdf.set_text_color(71, 85, 105)
        pdf.cell(0, 7, f"Tarih: {datetime.now().strftime('%d.%m.%Y')}",
                 new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        if ogretmen_adi:
            pdf.cell(0, 7, f"Hazırlayan: {ogretmen_adi}",
                     new_x=XPos.LMARGIN, new_y=YPos.NEXT)

        if secilen_basliklar:
            baslik_str = ", ".join(secilen_basliklar)
            pdf.cell(0, 7, f"Değerlendirme Alanları: {baslik_str}",
                     new_x=XPos.LMARGIN, new_y=YPos.NEXT)

        pdf.ln(5)

        # Özet metni
        pdf.set_font(font_adi, "", 10)
        pdf.set_text_color(30, 41, 59)

        satirlar = ozet_metni.split("\n")
        for satir in satirlar:
            satir = satir.strip()
            if not satir:
                pdf.ln(3)
                continue

            # Her satırdan önce X pozisyonunu sıfırla
            pdf.set_x(pdf.l_margin)

            if satir.startswith("##"):
                satir = satir.lstrip("#").strip()
                pdf.ln(3)
                pdf.set_font(font_adi, "B", 12)
                pdf.set_text_color(59, 130, 246)
                pdf.cell(0, 8, satir,
                         new_x=XPos.LMARGIN, new_y=YPos.NEXT)
                pdf.set_font(font_adi, "", 10)
                pdf.set_text_color(30, 41, 59)
            elif satir.startswith("**") and satir.endswith("**"):
                satir_temiz = satir.strip("*").strip()
                pdf.set_font(font_adi, "B", 11)
                pdf.cell(0, 7, satir_temiz,
                         new_x=XPos.LMARGIN, new_y=YPos.NEXT)
                pdf.set_font(font_adi, "", 10)
            elif satir.startswith("- ") or satir.startswith("\u2022 "):
                madde = satir.lstrip("- \u2022").strip()
                pdf.multi_cell(0, 6, f"    \u2022 {madde}",
                               new_x=XPos.LMARGIN, new_y=YPos.NEXT)
            else:
                pdf.multi_cell(0, 6, satir,
                               new_x=XPos.LMARGIN, new_y=YPos.NEXT)

        # Yasal uyarı
        pdf.ln(10)
        pdf.set_x(pdf.l_margin)
        pdf.set_font(font_adi, "", 7)
        pdf.set_text_color(148, 163, 184)
        pdf.multi_cell(0, 4,
                       "Bu belge eğitim amaçlı hazırlanmıştır. Klinik tanı yerine geçmez. "
                       "Profesyonel değerlendirme için uzman görüşü alınız.",
                       new_x=XPos.LMARGIN, new_y=YPos.NEXT)

        return pdf.output()
    except Exception:
        return None


def aile_docx_olustur(ogrenci_adi, ozet_metni, ogretmen_adi="", secilen_basliklar=None):
    """Aile bilgilendirme DOCX raporu oluştur"""
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()

        # Stil
        style = doc.styles["Normal"]
        style.font.name = "DejaVu Sans"
        style.font.size = Pt(11)

        # Başlık
        baslik = doc.add_heading("Veli Bilgilendirme Özeti", level=0)
        baslik.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in baslik.runs:
            run.font.color.rgb = RGBColor(30, 41, 59)

        alt = doc.add_paragraph("Eğitim Check-Up Platformu")
        alt.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in alt.runs:
            run.font.color.rgb = RGBColor(100, 116, 139)
            run.font.size = Pt(10)

        doc.add_paragraph("")

        # Bilgi alanı
        bilgi = doc.add_paragraph()
        bilgi.add_run("Öğrenci: ").bold = True
        bilgi.add_run(ogrenci_adi)

        tarih_p = doc.add_paragraph()
        tarih_p.add_run("Tarih: ").bold = True
        tarih_p.add_run(datetime.now().strftime("%d.%m.%Y"))

        if ogretmen_adi:
            ogr_p = doc.add_paragraph()
            ogr_p.add_run("Hazırlayan: ").bold = True
            ogr_p.add_run(ogretmen_adi)

        if secilen_basliklar:
            alan_p = doc.add_paragraph()
            alan_p.add_run("Değerlendirme Alanları: ").bold = True
            alan_p.add_run(", ".join(secilen_basliklar))

        doc.add_paragraph("")

        # Özet metni
        satirlar = ozet_metni.split("\n")
        for satir in satirlar:
            satir = satir.strip()
            if not satir:
                doc.add_paragraph("")
                continue

            if satir.startswith("##"):
                satir = satir.lstrip("#").strip()
                h = doc.add_heading(satir, level=2)
                for run in h.runs:
                    run.font.color.rgb = RGBColor(59, 130, 246)
            elif satir.startswith("**") and satir.endswith("**"):
                p = doc.add_paragraph()
                run = p.add_run(satir.strip("*").strip())
                run.bold = True
            elif satir.startswith("- ") or satir.startswith("\u2022 "):
                doc.add_paragraph(satir.lstrip("- \u2022").strip(), style="List Bullet")
            else:
                doc.add_paragraph(satir)

        # Yasal uyarı
        doc.add_paragraph("")
        uyari = doc.add_paragraph(
            "Bu belge eğitim amaçlı hazırlanmıştır. Klinik tanı yerine geçmez. "
            "Profesyonel değerlendirme için uzman görüşü alınız."
        )
        for run in uyari.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(148, 163, 184)

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf.getvalue()
    except Exception:
        return None

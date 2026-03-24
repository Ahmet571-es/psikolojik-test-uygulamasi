# -*- coding: utf-8 -*-
"""Eğitim Check-Up - PDF ve DOCX Rapor Üretim Modülü"""

import io
import os
from datetime import datetime
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np


# DejaVu Sans font ayarı (Türkçe karakter desteği)
def _font_ayarla():
    """Matplotlib ve FPDF için font ayarlarını yap"""
    font_yollari = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/dejavu/DejaVuSans.ttf",
        os.path.join(os.path.dirname(__file__), "fonts", "DejaVuSans.ttf"),
    ]

    font_yolu = None
    for yol in font_yollari:
        if os.path.exists(yol):
            font_yolu = yol
            break

    if font_yolu:
        import matplotlib.font_manager as fm
        fm.fontManager.addfont(font_yolu)
        plt.rcParams["font.family"] = "DejaVu Sans"
    else:
        plt.rcParams["font.family"] = "sans-serif"

    plt.rcParams["axes.unicode_minus"] = False
    return font_yolu

FONT_YOLU = _font_ayarla()


def radar_grafigi_olustur(etiketler, degerler, baslik=""):
    """Radar (örümcek ağı) grafiği oluştur"""
    try:
        etiketler = np.array(etiketler)
        degerler = np.array(degerler, dtype=float)

        acilar = np.linspace(0, 2 * np.pi, len(etiketler), endpoint=False).tolist()
        degerler_kapali = np.concatenate((degerler, [degerler[0]]))
        acilar += acilar[:1]

        fig, ax = plt.subplots(figsize=(6, 6), subplot_kw=dict(polar=True))
        ax.fill(acilar, degerler_kapali, color="#3B82F6", alpha=0.2)
        ax.plot(acilar, degerler_kapali, color="#2563EB", linewidth=2, marker="o", markersize=4)
        ax.set_yticklabels([])
        ax.set_xticks(acilar[:-1])
        ax.set_xticklabels(etiketler, fontsize=9, weight="bold")
        if baslik:
            ax.set_title(baslik, y=1.12, fontsize=13, fontweight="bold")
        ax.spines["polar"].set_visible(False)
        ax.grid(color="#E2E8F0", linestyle="--")
        plt.tight_layout()
        return fig
    except Exception:
        return None


def cubuk_grafigi_olustur(etiketler, degerler, baslik=""):
    """Yatay çubuk grafiği oluştur"""
    try:
        fig, ax = plt.subplots(figsize=(8, max(len(etiketler) * 0.5, 3)))
        renkler = plt.cm.Blues(np.linspace(0.4, 0.8, len(etiketler)))
        bars = ax.barh(etiketler, degerler, color=renkler, height=0.6)
        ax.set_xlim(0, 100)
        ax.set_xlabel("Yüzde (%)")
        if baslik:
            ax.set_title(baslik, fontsize=13, fontweight="bold")

        for bar, val in zip(bars, degerler):
            ax.text(bar.get_width() + 1, bar.get_y() + bar.get_height() / 2,
                    f"%{val}", va="center", fontsize=10)

        plt.tight_layout()
        return fig
    except Exception:
        return None


def _grafik_bytes(fig):
    """Matplotlib figürünü bytes'a dönüştür"""
    if fig is None:
        return None
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150, bbox_inches="tight")
    plt.close(fig)
    buf.seek(0)
    return buf


# --- PDF Üretimi ---

def pdf_olustur(ogrenci_adi, test_adi, sonuclar, grafik_verileri=None, ek_metin="", ogretmen_adi=""):
    """PDF rapor oluştur"""
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    pdf = FPDF()
    font_adi = "Helvetica"

    if FONT_YOLU and os.path.exists(FONT_YOLU):
        pdf.add_font("DejaVu", "", FONT_YOLU)
        bold_yol = FONT_YOLU.replace("DejaVuSans.ttf", "DejaVuSans-Bold.ttf")
        if os.path.exists(bold_yol):
            pdf.add_font("DejaVu", "B", bold_yol)
        else:
            pdf.add_font("DejaVu", "B", FONT_YOLU)
        font_adi = "DejaVu"

    pdf.add_page()

    # Başlık
    pdf.set_font(font_adi, "B", 16)
    pdf.set_text_color(30, 41, 59)
    pdf.cell(0, 10, "Eğitim Check-Up",
             new_x=XPos.LMARGIN, new_y=YPos.NEXT, align="C")
    pdf.set_font(font_adi, "", 9)
    pdf.set_text_color(100, 116, 139)
    pdf.cell(0, 6, "Eğitim Psikolojisi Değerlendirme Raporu",
             new_x=XPos.LMARGIN, new_y=YPos.NEXT, align="C")
    pdf.line(10, pdf.get_y() + 2, 200, pdf.get_y() + 2)
    pdf.ln(8)

    # Öğrenci ve test bilgileri
    pdf.set_font(font_adi, "B", 14)
    pdf.set_text_color(30, 41, 59)
    pdf.cell(0, 10, test_adi, new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(2)

    pdf.set_font(font_adi, "", 11)
    pdf.set_text_color(51, 65, 85)
    pdf.cell(0, 7, f"Öğrenci: {ogrenci_adi}",
             new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.cell(0, 7, f"Tarih: {datetime.now().strftime('%d.%m.%Y %H:%M')}",
             new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    if ogretmen_adi:
        pdf.cell(0, 7, f"Öğretmen: {ogretmen_adi}",
                 new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(5)

    # Sonuçlar
    if isinstance(sonuclar, dict):
        for anahtar, deger in sonuclar.items():
            if isinstance(deger, (list, dict)):
                continue
            pdf.set_font(font_adi, "B", 11)
            baslik_text = anahtar.replace("_", " ").title()
            pdf.cell(60, 8, f"{baslik_text}:")
            pdf.set_font(font_adi, "", 11)
            pdf.cell(0, 8, str(deger), new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    # Grafik ekleme
    if grafik_verileri:
        pdf.ln(5)
        for grafik_bilgi in grafik_verileri:
            etiketler = grafik_bilgi.get("etiketler", [])
            degerler = grafik_bilgi.get("degerler", [])
            baslik_g = grafik_bilgi.get("baslik", "")
            tip = grafik_bilgi.get("tip", "cubuk")

            if tip == "radar":
                fig = radar_grafigi_olustur(etiketler, degerler, baslik_g)
            else:
                fig = cubuk_grafigi_olustur(etiketler, degerler, baslik_g)

            buf = _grafik_bytes(fig)
            if buf:
                import tempfile
                with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
                    tmp.write(buf.getvalue())
                    tmp_yol = tmp.name

                if pdf.get_y() > 180:
                    pdf.add_page()
                pdf.image(tmp_yol, x=30, w=150)
                pdf.ln(5)
                os.unlink(tmp_yol)

    # Ek metin
    if ek_metin:
        if pdf.get_y() > 230:
            pdf.add_page()
        pdf.ln(5)
        pdf.set_x(pdf.l_margin)
        pdf.set_font(font_adi, "", 10)
        pdf.multi_cell(0, 6, ek_metin, new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    # Yasal uyarı
    pdf.ln(10)
    pdf.set_x(pdf.l_margin)
    pdf.set_font(font_adi, "", 8)
    pdf.set_text_color(128, 128, 128)
    pdf.multi_cell(0, 5, "Bu rapor eğitim amaçlı hazırlanmıştır. Klinik tanı yerine geçmez. "
                         "Profesyonel değerlendirme için uzman görüşü alınız.",
                   new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    return pdf.output()


# --- DOCX Üretimi ---

def docx_olustur(ogrenci_adi, test_adi, sonuclar, grafik_verileri=None, ek_metin="", ogretmen_adi=""):
    """DOCX rapor oluştur"""
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Stil ayarları
    style = doc.styles["Normal"]
    font = style.font
    font.name = "DejaVu Sans"
    font.size = Pt(11)

    # Başlık
    baslik = doc.add_heading("Eğitim Check-Up", level=0)
    baslik.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in baslik.runs:
        run.font.color.rgb = RGBColor(30, 41, 59)

    alt_baslik = doc.add_paragraph("Eğitim Psikolojisi Değerlendirme Raporu")
    alt_baslik.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("")

    # Test bilgileri
    doc.add_heading(test_adi, level=1)
    doc.add_paragraph(f"Öğrenci: {ogrenci_adi}")
    doc.add_paragraph(f"Tarih: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
    if ogretmen_adi:
        doc.add_paragraph(f"Öğretmen: {ogretmen_adi}")

    doc.add_paragraph("")

    # Sonuçlar tablosu
    if isinstance(sonuclar, dict):
        doc.add_heading("Sonuçlar", level=2)
        tablo = doc.add_table(rows=1, cols=2)
        tablo.style = "Light Grid Accent 1"
        hdr = tablo.rows[0].cells
        hdr[0].text = "Ölçüt"
        hdr[1].text = "Değer"

        for anahtar, deger in sonuclar.items():
            if isinstance(deger, (list, dict)):
                continue
            satir = tablo.add_row().cells
            satir[0].text = anahtar.replace("_", " ").title()
            satir[1].text = str(deger)

    # Grafik ekleme
    if grafik_verileri:
        doc.add_paragraph("")
        doc.add_heading("Grafikler", level=2)
        for grafik_bilgi in grafik_verileri:
            etiketler = grafik_bilgi.get("etiketler", [])
            degerler = grafik_bilgi.get("degerler", [])
            baslik_text = grafik_bilgi.get("baslik", "")
            tip = grafik_bilgi.get("tip", "cubuk")

            if tip == "radar":
                fig = radar_grafigi_olustur(etiketler, degerler, baslik_text)
            else:
                fig = cubuk_grafigi_olustur(etiketler, degerler, baslik_text)

            buf = _grafik_bytes(fig)
            if buf:
                doc.add_picture(buf, width=Inches(5.5))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Ek metin
    if ek_metin:
        doc.add_paragraph("")
        doc.add_heading("Değerlendirme", level=2)
        doc.add_paragraph(ek_metin)

    # Yasal uyarı
    doc.add_paragraph("")
    uyari = doc.add_paragraph(
        "Bu rapor eğitim amaçlı hazırlanmıştır. Klinik tanı yerine geçmez. "
        "Profesyonel değerlendirme için uzman görüşü alınız."
    )
    for run in uyari.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(128, 128, 128)

    # BytesIO'ya kaydet
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()
